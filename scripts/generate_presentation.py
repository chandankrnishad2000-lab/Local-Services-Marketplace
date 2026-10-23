from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

OUTPUT = Path(__file__).resolve().parents[1] / "LocalPulse_Project_Presentation.pptx"

AUTHOR_NAME = "Chandan Kumar"
INSTITUTE = "LPU"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TITLE_COLOR = RGBColor(30, 64, 175)
ACCENT = RGBColor(236, 91, 19)
TEXT = RGBColor(55, 65, 81)


def set_background(slide, color=(248, 250, 252)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)


def add_title_slide(
    title,
    subtitle,
    footer="LocalPulse | Local Services Marketplace",
    presented_by=None,
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)

    shape = slide.shapes.add_textbox(Inches(0.7), Inches(0.9), Inches(11.9), Inches(1.2))
    tf = shape.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR

    shape = slide.shapes.add_textbox(Inches(0.72), Inches(2.0), Inches(11.5), Inches(0.8))
    tf = shape.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = subtitle
    r.font.size = Pt(18)
    r.font.color.rgb = TEXT

    if presented_by:
        shape = slide.shapes.add_textbox(Inches(0.72), Inches(2.6), Inches(11.5), Inches(0.5))
        tf = shape.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = presented_by
        r.font.size = Pt(14)
        r.font.color.rgb = TEXT

    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.72), Inches(3.0), Inches(4.1), Inches(0.12))
    box.fill.solid()
    box.fill.fore_color.rgb = ACCENT
    box.line.color.rgb = ACCENT

    shape = slide.shapes.add_textbox(Inches(0.72), Inches(6.7), Inches(11.5), Inches(0.3))
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = footer
    r.font.size = Pt(12)
    r.font.color.rgb = TEXT


def add_bullets_slide(title, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)

    header = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.7))
    tf = header.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.72), Inches(1.1), Inches(11.5), Inches(0.4))
        tf = sub.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        r.font.size = Pt(14)
        r.font.color.rgb = TEXT

    body = slide.shapes.add_textbox(Inches(0.85), Inches(1.65), Inches(11.7), Inches(5.4))
    tf = body.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)

    first = True
    for bullet in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = f"• {bullet}"
        p.level = 0
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT
        p.space_after = Pt(10)


def add_two_column_slide(title, left_title, left_points, right_title, right_points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)

    header = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.7))
    tf = header.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR

    for x, heading, points in [
        (0.8, left_title, left_points),
        (6.95, right_title, right_points),
    ]:
        shape = slide.shapes.add_textbox(Inches(x), Inches(1.35), Inches(5.6), Inches(0.4))
        tf = shape.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = heading
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = ACCENT

        body = slide.shapes.add_textbox(Inches(x), Inches(1.8), Inches(5.6), Inches(4.8))
        tf = body.text_frame
        tf.word_wrap = True
        first = True
        for point in points:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = f"• {point}"
            p.font.size = Pt(16)
            p.font.color.rgb = TEXT
            p.space_after = Pt(8)


add_title_slide(
    "LocalPulse",
    "A modern local services marketplace for customers, local pros, and admins",
    presented_by=f"Presented by: {AUTHOR_NAME} ({INSTITUTE})",
)

add_bullets_slide(
    "Project Goal",
    [
        "Connect people who need local services with trusted service providers in their city.",
        "Support discovery, booking, payments, reviews, and communication in one platform.",
        "Keep the current app as a demo-friendly marketplace with realistic workflows and dummy integrations."
    ],
    "This project is designed as a full-stack marketplace experience."
)

add_two_column_slide(
    "User Roles",
    "Customer",
    [
        "Search and browse listings",
        "Book services and add tips",
        "Leave reviews after service completion",
        "Message the provider through a booking thread"
    ],
    "Local Pro / Admin",
    [
        "Local pros create and manage listings",
        "Connect dummy Stripe onboarding",
        "Admins review reports and platform activity",
        "Role-based access controls protect each area"
    ]
)

add_bullets_slide(
    "Main User Flow",
    [
        "Customer signs up or logs in using password or OTP.",
        "Customer browses services by category, city, and price.",
        "Customer opens a listing and requests a booking.",
        "Booking is created, payment record is generated, and checkout is simulated.",
        "Customer and local pro can continue communication and reviews afterward."
    ]
)

add_bullets_slide(
    "Core Marketplace Features",
    [
        "Homepage with hero, features, services, cities, testimonials, and CTA sections.",
        "Listing discovery page with filters and detail view.",
        "Booking flow with booking summary, tip selection, and payment placeholder.",
        "Profile page, admin moderation page, and local pro dashboard.",
        "Contact page with email and phone details."
    ]
)

add_two_column_slide(
    "Technology Stack",
    "Frontend",
    [
        "Next.js App Router",
        "React + TypeScript",
        "Bootstrap + Tailwind styling",
        "Client and server components"
    ],
    "Backend",
    [
        "Plain PHP backend in `backend/`",
        "SQLite database",
        "PHP sessions for auth",
        "Dummy Stripe and OTP flows"
    ]
)

add_bullets_slide(
    "Architecture Overview",
    [
        "Next.js handles the UI and page routing.",
        "PHP backend exposes API endpoints for auth, listings, bookings, reviews, messages, and admin reports.",
        "Frontend fetches PHP API data using shared client/server helpers.",
        "Old Next.js API and Prisma backend were removed from runtime paths.",
        "This keeps the stack simple, local, and demo-friendly."
    ]
)

add_two_column_slide(
    "Backend API (PHP)",
    "Core Routes",
    [
        "Auth: register, login, logout, me",
        "Listings: list, create, details, update, delete",
        "Bookings: create, list, update status",
    ],
    "Supporting Routes",
    [
        "Messages: thread chat between customer and provider",
        "Reviews: post-review after booking",
        "Reports: admin moderation dashboard",
        "Stripe (dummy): checkout/connect/webhook",
        "OTP (dummy): request/verify",
    ],
)

add_bullets_slide(
    "Database Schema (SQLite)",
    [
        "users: accounts with role (CUSTOMER / LOCAL_PRO / ADMIN)",
        "profiles: basic profile fields (name, phone, etc.)",
        "listings: service title, description, city, category, price",
        "bookings: customer ↔ listing, dates/status, tip",
        "payments: stored payment/checkout metadata (dummy in demo)",
        "threads + messages: booking chat between users",
        "reviews: rating + comment linked to booking/listing",
        "reports: admin-side report records",
        "otp_codes: temporary OTP codes for demo login",
    ],
    "Tables are auto-created by the backend for easy local setup."
)

add_bullets_slide(
    "Dummy Integrations",
    [
        "Stripe checkout is simulated instead of using live payment keys.",
        "OTP requests return the code for easy testing.",
        "Sessions are handled by PHP cookies.",
        "The project is intentionally lightweight for presentation and demo use."
    ]
)

add_bullets_slide(
    "How To Run (Demo)",
    [
        "Start backend: php -S localhost:8000 -t backend/public",
        "Start frontend: npm run dev (opens http://localhost:3000)",
        "Register a customer and a local pro.",
        "Local pro creates a listing.",
        "Customer books the listing and completes dummy checkout.",
        "Leave a review and send messages from booking thread.",
    ],
    "Two-process local setup: Next.js UI + PHP API."
)

add_bullets_slide(
    "Why This Project Matters",
    [
        "Solves a real local-services discovery and booking problem.",
        "Shows role-based product design and marketplace thinking.",
        "Demonstrates full-stack integration across frontend and backend.",
        "Is easy to present because the flows are concrete and visually clear."
    ]
)

add_bullets_slide(
    "Future Scope",
    [
        "Replace dummy Stripe with live payments + refunds + payout tracking.",
        "Real OTP delivery via SMS/Email and password reset emails.",
        "Advanced search: category/city filters, sorting, pagination.",
        "Provider availability calendar and reschedule/cancel rules.",
        "Notifications: email/SMS/in-app for booking updates.",
        "Admin tools: user verification, dispute handling, content moderation.",
        "Performance/security hardening: rate-limits, CSRF strategy, audit logs.",
    ]
)

add_title_slide(
    "Thank You",
    "LocalPulse — your local services marketplace demo",
    presented_by=f"{AUTHOR_NAME} | {INSTITUTE}",
)

prs.save(str(OUTPUT))
print(f"Created {OUTPUT}")
